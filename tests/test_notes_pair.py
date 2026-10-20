from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import render_notes_pair as pair


def sample():
    return json.loads((ROOT / 'examples/notes-data.json').read_text(encoding='utf-8'))


class PairTests(unittest.TestCase):
    def test_both_files_preserve_all_learning_content(self):
        data = sample()
        data['sources'] = [{'label': '核实来源', 'url': 'https://example.com/reference'}]
        data['sections'][0]['paragraphs'] = ['条件：时间间隔为正。']
        with tempfile.TemporaryDirectory() as tmp:
            html, word = pair.export_pair(data, Path(tmp) / '笔记')
            self.assertGreater(html.stat().st_size, 0)
            Document(word)  # A real, readable Word package, not renamed HTML.
            with ZipFile(word) as z:
                root = ET.fromstring(z.read('word/document.xml'))
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                text = '\n'.join(root.itertext())
                for needle in [data['title'], data['summary'], '条件：时间间隔为正。',
                               data['questions'][0]['answer'], data['appendices'][0]['title'],
                               'https://example.com/reference']:
                    self.assertIn(needle, text)
                    self.assertIn(needle, html.read_text(encoding='utf-8'))
                self.assertEqual(len(root.findall('.//w:tbl', ns)), 2)
                self.assertEqual(len(root.findall('.//w:tblHeader', ns)), 2)
                for row in data['sections'][0]['table']['rows']:
                    for cell in row:
                        self.assertIn(cell, text)

    def test_questions_are_separated_from_answers(self):
        doc = Document(BytesIO(pair.build_docx(sample())))
        texts = [p.text for p in doc.paragraphs]
        self.assertLess(texts.index('自测题'), texts.index('参考答案'))
        self.assertLess(next(i for i, t in enumerate(texts) if sample()['questions'][1]['question'] in t),
                        texts.index('参考答案'))

    def test_existing_either_file_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp) / 'notes'
            base.with_suffix('.docx').write_bytes(b'original')
            with self.assertRaises(FileExistsError):
                pair.export_pair(sample(), base)
            self.assertEqual(base.with_suffix('.docx').read_bytes(), b'original')
            self.assertFalse(base.with_suffix('.html').exists())

    def test_word_generation_failure_leaves_no_partial_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(pair, 'build_docx', side_effect=RuntimeError('Word failed')):
                with self.assertRaises(RuntimeError):
                    pair.export_pair(sample(), Path(tmp) / 'notes')
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_second_write_failure_rolls_back_first_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            original_link = pair.os.link
            calls = []
            def link(source, target):
                calls.append(target)
                if len(calls) == 2:
                    raise OSError('cannot write Word')
                original_link(source, target)
            with patch.object(pair.os, 'link', side_effect=link):
                with self.assertRaises(OSError):
                    pair.export_pair(sample(), Path(tmp) / 'notes')
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_overwrite_failure_restores_previous_pair(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp) / 'notes'
            for suffix in ('.html', '.docx'):
                base.with_suffix(suffix).write_bytes(b'previous ' + suffix.encode())
            original_replace = pair.os.replace
            calls = []
            def replace(source, target):
                calls.append(target)
                if len(calls) == 2:
                    raise OSError('cannot replace Word')
                original_replace(source, target)
            with patch.object(pair.os, 'replace', side_effect=replace):
                with self.assertRaises(OSError):
                    pair.export_pair(sample(), base, overwrite=True)
            for suffix in ('.html', '.docx'):
                self.assertEqual(base.with_suffix(suffix).read_bytes(), b'previous ' + suffix.encode())


if __name__ == '__main__':
    unittest.main()
