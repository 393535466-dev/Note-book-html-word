#!/usr/bin/env python3
"""Export one notes JSON as matching HTML and editable Word documents."""
import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile

from render_notes import render_notes


def build_docx(data, cjk_font=None):
    """Use the same validated content as HTML; Word expands all appendices."""
    render_notes(data)
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError('Word export requires python-docx. Use a runtime with python-docx installed.') from exc

    cjk_font = cjk_font or ('Hiragino Sans GB' if sys.platform == 'darwin' else
                            'Microsoft YaHei' if sys.platform == 'win32' else 'Noto Sans CJK SC')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(21)
    width_mm = 168

    for name, size, color in [('Normal', 10.5, '182338'), ('Title', 22, '000000'),
                              ('Heading 1', 14, '182338'), ('Heading 2', 11.5, '182338')]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), cjk_font)
        # Override template theme fonts and inherited title border decoration.
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        ppr = style.element.get_or_add_pPr()
        old_border = ppr.find(qn('w:pBdr'))
        if old_border is not None:
            ppr.remove(old_border)
        if name == 'Title':
            borders = OxmlElement('w:pBdr')
            for side in ('top', 'bottom', 'left', 'right', 'between'):
                border = OxmlElement('w:' + side)
                border.set(qn('w:val'), 'nil')
                borders.append(border)
            ppr.append(borders)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25
        if name != 'Normal':
            style.paragraph_format.keep_with_next = True
    doc.styles['Normal'].paragraph_format.widow_control = True
    doc.core_properties.title = data['title']
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.add_paragraph(data['title'], 'Title')
    if data.get('meta'):
        p = doc.add_paragraph(data['meta'])
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor.from_string('607086')
    doc.add_paragraph(data['summary'])

    def list_items(items, ordered=False):
        for index, text in enumerate(items, 1):
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.left_indent = Mm(4)
            p.paragraph_format.first_line_indent = Mm(-4)
            p.add_run(f'{index}. ' if ordered else '• ')
            p.add_run(text)

    def add_table(value):
        columns = value['columns']
        count = len(columns)
        # Compact status/level columns leave space for explanations and examples.
        weights = [0.65 if c in ('级别', '状态', '负责人', '截止日') else
                   1.0 if c in ('知识点', '任务', '术语', '内容') else 2.0 for c in columns]
        widths = [width_mm * w / sum(weights) for w in weights]
        cap = doc.add_paragraph(value['caption'])
        cap.paragraph_format.space_before = Pt(5)
        cap.paragraph_format.keep_with_next = True
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.color.rgb = RGBColor.from_string('607086')
        table = doc.add_table(rows=1, cols=count)
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Mm(width)
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement(f'w:{side}')
            for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                border.set(qn('w:' + key), val)
            borders.append(border)
        props.append(borders)
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for values in [columns, *value['rows']]:
            row = table.rows[0] if values is columns else table.add_row()
            index = len(table.rows) - 1
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for col, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Mm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = text
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'left', 'bottom', 'right'):
                    edge = OxmlElement('w:' + side)
                    edge.set(qn('w:w'), '90')
                    edge.set(qn('w:type'), 'dxa')
                    margins.append(edge)
                tcpr.append(margins)
                if values is columns or index % 2 == 0:
                    fill = OxmlElement('w:shd')
                    fill.set(qn('w:fill'), 'EDF2FA' if values is columns else 'F8FAFD')
                    tcpr.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = False
                    p.paragraph_format.line_spacing = 1.2
                    if columns[col] in ('级别', '状态', '负责人', '截止日'):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.bold = values is columns
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.space_before = Pt(2)
        spacer.paragraph_format.line_spacing = Pt(6)
        spacer.add_run().font.size = Pt(6)

    def add_block(block):
        if block.get('intro'):
            doc.add_paragraph(block['intro'])
        for text in block.get('paragraphs', []):
            doc.add_paragraph(text)
        list_items(block.get('bullets', []))
        list_items(block.get('steps', []), ordered=True)
        if 'table' in block:
            add_table(block['table'])

    if data.get('highlights'):
        doc.add_heading('重点速览', 1)
        list_items(data['highlights'])
    for block in data['sections']:
        doc.add_heading(block['title'], 1)
        add_block(block)
    if data.get('questions'):
        doc.add_heading('自测题', 1)
        list_items([q['question'] for q in data['questions']], ordered=True)
        doc.add_heading('参考答案', 1)
        for i, q in enumerate(data['questions'], 1):
            doc.add_paragraph(f'{i}. {q["answer"]}')
    for block in data.get('appendices', []):
        doc.add_heading(block['title'], 1)
        add_block(block)
    if data.get('sources'):
        doc.add_heading('外部核实来源', 1)
        for source in data['sources']:
            p = doc.add_paragraph(source['label'] + '\n')
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(source['url'], RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            text = OxmlElement('w:t')
            text.text = source['url']
            run.append(text)
            link.append(run)
            p._p.append(link)
    if data.get('review'):
        doc.add_paragraph(data['review'])
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def export_pair(data, output_base, overwrite=False, cjk_font=None):
    """Generate both before publishing either; restore originals on write failure."""
    output_base = Path(output_base)
    paths = [Path(str(output_base) + '.html'), Path(str(output_base) + '.docx')]
    if not overwrite and any(path.exists() for path in paths):
        raise FileExistsError('Output exists. Choose a new basename or use --overwrite.')
    payloads = [render_notes(data).encode('utf-8'), build_docx(data, cjk_font=cjk_font)]
    output_base.parent.mkdir(parents=True, exist_ok=True)
    previous = [path.read_bytes() if path.exists() else None for path in paths]
    written = []
    with tempfile.TemporaryDirectory(prefix='.notes-', dir=output_base.parent) as folder:
        staged = []
        for i, payload in enumerate(payloads):
            path = Path(folder) / str(i)
            path.write_bytes(payload)
            staged.append(path)
        try:
            for source, target in zip(staged, paths):
                if overwrite:
                    os.replace(source, target)
                else:
                    os.link(source, target)
                written.append(target)
        except OSError:
            for target in written:
                backup = previous[paths.index(target)]
                if backup is None:
                    target.unlink()
                else:
                    target.write_bytes(backup)
            raise
    return paths


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output_base', type=Path, help='Output basename without .html/.docx')
    parser.add_argument('--overwrite', action='store_true')
    parser.add_argument('--cjk-font', help='Installed Chinese font family for Word')
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding='utf-8'))
        paths = export_pair(data, args.output_base, args.overwrite, args.cjk_font)
    except (OSError, ValueError, RuntimeError) as exc:
        parser.exit(1, f'Cannot export notes: {exc}\n')
    for path in paths:
        print(f'Created {path.resolve()}')


if __name__ == '__main__':
    main()
